from pathlib import Path

from langchain_core.tools import tool


@tool
def list_doc_sections(file_path: str) -> str:
	"""List all headings/sections in a Word document (.docx) to get a structural overview before deciding what to read."""
	try:
		from docx import Document
		doc = Document(file_path)
		headings = []
		for para in doc.paragraphs:
			if para.style.name.startswith("Heading"):
				level = para.style.name.replace("Heading ", "")
				indent = "  " * (int(level) - 1) if level.isdigit() else ""
				headings.append(f"{indent}- [{para.style.name}] {para.text}")
		if not headings:
			paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
			return (
				f"'{Path(file_path).name}' has no headings. {len(paras)} paragraphs total.\n"
				f"First paragraph: {paras[0][:200] if paras else 'empty'}"
			)
		return f"Sections in '{Path(file_path).name}':\n" + "\n".join(headings)
	except Exception as e:
		return f"Error reading '{file_path}': {e}"


@tool
def read_doc_section(file_path: str, heading: str) -> str:
	"""Read the content under a specific heading in a Word document. Only call for sections relevant to the user's objective."""
	try:
		from docx import Document
		doc = Document(file_path)
		collecting = False
		content = []

		for para in doc.paragraphs:
			is_heading = para.style.name.startswith("Heading")

			if is_heading and heading.lower() in para.text.lower():
				collecting = True
				content.append(f"## {para.text}")
				continue

			if collecting:
				if is_heading:
					break  # hit the next section — stop
				if para.text.strip():
					content.append(para.text.strip())

		if not content:
			return f"Section '{heading}' not found in '{Path(file_path).name}'."
		return "\n".join(content)
	except Exception as e:
		return f"Error reading section '{heading}': {e}"
